"""Change audit, exports, alert deduplication and the emailed reports.

These cover the behaviours that are expensive to get wrong and invisible when they
break: an audit trail that forgets who acted, an export that silently truncates,
and an alerting system that either floods inboxes or stays quiet during an outage.
"""
from __future__ import annotations

from datetime import timedelta
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core import mail
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from c360 import changes
from c360.models import AlertState, AuditEvent, MetricMinute, Profile
from c360.reports import alerts, datasets, digest, errorlog, tabular

LOCMEM = 'django.core.mail.backends.locmem.EmailBackend'


class ChangeAuditTests(TestCase):
    """The history tables, and the feed that reads them back."""

    def setUp(self):
        self.actor = User.objects.create_user('acting.admin', password='x')
        self.target = User.objects.create_user('target.rm', password='x')

    def _as(self, user):
        """simple_history reads the acting user off the thread-local request."""
        from simple_history.models import HistoricalRecords

        class _Req:
            pass
        req = _Req()
        req.user = user
        HistoricalRecords.context.request = req
        return req

    def tearDown(self):
        from simple_history.models import HistoricalRecords
        try:
            del HistoricalRecords.context.request
        except AttributeError:
            pass

    def test_password_hash_is_not_kept_in_history(self):
        """Historical password hashes are crackable; the table must not hold them."""
        fields = {f.name for f in User.history.model._meta.fields}
        self.assertNotIn('password', fields)

    def test_update_records_actor_and_field_diff(self):
        self._as(self.actor)
        self.target.is_staff = True
        self.target.save()

        rows = changes.feed(limit=20, model='user')
        updates = [r for r in rows if r['action'] == 'updated' and r['object_label'] == 'target.rm']
        self.assertTrue(updates, 'the update was not recorded')
        row = updates[0]
        self.assertEqual(row['username'], 'acting.admin')
        self.assertFalse(row['external_actor'])
        diff = {c['field']: (c['old'], c['new']) for c in row['changes']}
        self.assertEqual(diff['is_staff'], ('no', 'yes'))

    def test_profile_allocation_change_is_recorded(self):
        self._as(self.actor)
        profile = Profile.objects.get(user=self.target)
        profile.sales_code = 'SC-4321'
        profile.save()

        rows = changes.feed(limit=20, model='profile')
        diff = {c['field']: c['new'] for r in rows for c in r['changes']}
        self.assertEqual(diff.get('sales_code'), 'SC-4321')

    def test_role_change_renders_group_names_not_row_ids(self):
        """An m2m diff is through-table rows; the feed must show what a person means."""
        self._as(self.actor)
        group = Group.objects.create(name='c360_management')
        self.target.groups.add(group)

        rows = changes.feed(limit=20, model='user')
        renders = [c['new'] for r in rows for c in r['changes'] if c['field'] == 'groups']
        self.assertTrue(renders, 'the role change was not recorded')
        self.assertIn('c360_management', renders[0])
        self.assertNotIn('{', renders[0])

    def test_sso_actor_without_local_account_is_still_named(self):
        """A portfolio SSO caller has no local row to hang the FK on. The audit must
        still say who it was, rather than reporting the change as anonymous."""
        class _ClaimsLike:
            is_authenticated = True
            username = 'portfolio.only'
        self._as(_ClaimsLike())

        self.target.first_name = 'Renamed'
        self.target.save()

        rows = [r for r in changes.feed(limit=20, model='user') if r['action'] == 'updated']
        self.assertTrue(rows)
        self.assertEqual(rows[0]['username'], 'portfolio.only')
        self.assertTrue(rows[0]['external_actor'])
        # And filtering by that name finds it.
        self.assertTrue(changes.feed(limit=20, username='portfolio.only'))

    def test_no_op_update_is_labelled_not_shown_as_empty(self):
        self._as(self.actor)
        self.target.save()                     # nothing actually changed
        rows = [r for r in changes.feed(limit=20, model='user') if r['action'] == 'updated']
        self.assertTrue(rows)
        self.assertTrue(rows[0]['no_op'])
        self.assertEqual(rows[0]['changes'], [])

    def test_summary_counts_by_action(self):
        self._as(self.actor)
        self.target.is_active = False
        self.target.save()
        summary = changes.summary()
        self.assertGreaterEqual(summary['by_action']['created'], 2)   # the two users
        self.assertGreaterEqual(summary['by_action']['updated'], 1)
        self.assertGreaterEqual(summary['actors'], 1)


class ExportEndpointTests(TestCase):
    def setUp(self):
        self.c = APIClient()
        AuditEvent.objects.create(kind='api', method='GET', route='/api/customers/',
                                  status=200, duration_ms=42, username='rm.one')
        AuditEvent.objects.create(kind='api', method='GET', route='/api/boom/',
                                  status=500, duration_ms=900, username='rm.two')

    def test_requires_admin(self):
        r = self.c.get('/api/observability/export/?dataset=activity')
        self.assertEqual(r.status_code, 403)

    def test_unknown_dataset_is_rejected_with_the_valid_names(self):
        r = self.c.get('/api/observability/export/?dataset=whatever', HTTP_X_C360_ADMIN='1')
        self.assertEqual(r.status_code, 400)
        self.assertIn('activity', r.json()['error']['detail'])

    def test_bad_format_is_rejected(self):
        r = self.c.get('/api/observability/export/?dataset=activity&fmt=exe',
                       HTTP_X_C360_ADMIN='1')
        self.assertEqual(r.status_code, 400)

    def test_csv_has_bom_header_row_and_every_row(self):
        r = self.c.get('/api/observability/export/?dataset=activity&fmt=csv',
                       HTTP_X_C360_ADMIN='1')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r['Content-Type'].startswith('text/csv'))
        self.assertIn('attachment; filename="c360-activity-trail-', r['Content-Disposition'])
        body = r.content
        # Excel needs the BOM to read UTF-8; without it non-ASCII names are mangled.
        self.assertTrue(body.startswith(b'\xef\xbb\xbf'))
        text = body.decode('utf-8-sig')
        lines = [ln for ln in text.splitlines() if ln.strip()]
        self.assertEqual(lines[0].split(',')[0], 'Time')
        self.assertEqual(len(lines), 3)                    # header + the two events
        self.assertEqual(r['X-C360-Row-Count'], '2')

    def test_xlsx_is_a_real_workbook(self):
        import io
        from openpyxl import load_workbook
        r = self.c.get('/api/observability/export/?dataset=errors&fmt=xlsx',
                       HTTP_X_C360_ADMIN='1')
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws['A1'].value, 'Error log')
        self.assertEqual(ws['A4'].value, 'Time')           # header row, under the title block
        self.assertEqual(ws.freeze_panes, 'A5')

    def test_every_registered_dataset_exports(self):
        for key in datasets.EXPORTABLE:
            with self.subTest(dataset=key):
                r = self.c.get(f'/api/observability/export/?dataset={key}&fmt=xlsx',
                               HTTP_X_C360_ADMIN='1')
                self.assertEqual(r.status_code, 200)


class ChangeFeedApiTests(TestCase):
    def setUp(self):
        self.c = APIClient()

    def test_requires_admin(self):
        self.assertEqual(self.c.get('/api/observability/changes/').status_code, 403)

    def test_returns_feed_summary_and_model_list(self):
        User.objects.create_user('someone', password='x')
        r = self.c.get('/api/observability/changes/', HTTP_X_C360_ADMIN='1')
        self.assertEqual(r.status_code, 200)
        payload = r.json()
        self.assertIn('results', payload)
        self.assertIn('summary', payload)
        self.assertIn('User account', [m['label'] for m in payload['models']])


class TabularTests(TestCase):
    def test_numbers_stay_numeric_and_none_becomes_blank(self):
        dataset = {
            'key': 'x', 'title': 'X', 'subtitle': '', 'generated_at': '', 'meta': {},
            'columns': [{'key': 'n', 'label': 'N', 'type': 'num'},
                        {'key': 't', 'label': 'T', 'type': 'text'}],
            'rows': [{'n': '42', 't': None}, {'n': 7, 't': 'ok'}],
            'row_count': 2,
        }
        text = tabular.to_csv(dataset).decode('utf-8-sig')
        rows = [ln for ln in text.splitlines() if ln.strip()]
        self.assertEqual(rows[1], '42,')                   # coerced to a number, blank text
        self.assertEqual(rows[2], '7,ok')

    def test_sheet_name_is_sanitised(self):
        self.assertEqual(tabular._sheet_name('a/b:c[d]' + 'x' * 40)[:7], 'a-b-c-d')
        self.assertLessEqual(len(tabular._sheet_name('y' * 60)), 31)


@override_settings(EMAIL_BACKEND=LOCMEM, C360_REPORT_RECIPIENTS=['ops@example.com'],
                   C360_REPORT_FROM='Reports <reports@example.com>')
class AlertTests(TestCase):
    FIRING = [{'key': 'warehouse_conn', 'severity': 'critical',
               'title': 'Warehouse unreachable', 'detail': 'auth rejected', 'hint': ''}]

    def setUp(self):
        mail.outbox = []
        AlertState.objects.all().delete()

    def test_incident_sends_once_then_stays_quiet(self):
        """A five-minute cron against an all-night outage must not send 100 emails."""
        with patch.object(alerts, 'evaluate', return_value=self.FIRING):
            first = alerts.run()
            for _ in range(5):
                alerts.run()
        self.assertEqual(first['sent'], 1)
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('Warehouse unreachable', mail.outbox[0].subject)
        self.assertEqual(mail.outbox[0].from_email, 'Reports <reports@example.com>')
        self.assertEqual(mail.outbox[0].to, ['ops@example.com'])

    def test_recovery_sends_exactly_one_resolved_mail(self):
        with patch.object(alerts, 'evaluate', return_value=self.FIRING):
            alerts.run()
        with patch.object(alerts, 'evaluate', return_value=[]):
            resolved = alerts.run()
            alerts.run()
        self.assertEqual(len(resolved['resolved']), 1)
        self.assertEqual(len(mail.outbox), 2)
        self.assertIn('Resolved', mail.outbox[1].subject)
        self.assertFalse(AlertState.objects.get(key='warehouse_conn').firing)

    def test_renotifies_once_the_quiet_interval_has_passed(self):
        with patch.object(alerts, 'evaluate', return_value=self.FIRING):
            alerts.run()
            state = AlertState.objects.get(key='warehouse_conn')
            state.last_notified_at = timezone.now() - timedelta(hours=24)
            state.save()
            again = alerts.run()
        self.assertEqual(len(again['ongoing']), 1)
        self.assertEqual(len(mail.outbox), 2)

    def test_dry_run_sends_nothing_and_records_nothing(self):
        with patch.object(alerts, 'evaluate', return_value=self.FIRING):
            alerts.run(dry_run=True)
        self.assertEqual(len(mail.outbox), 0)
        self.assertFalse(AlertState.objects.exists())

    def test_low_traffic_does_not_trip_the_error_rate_check(self):
        """One failure out of three requests is 33% — and means nothing."""
        self.assertEqual(alerts._check_error_rate(
            {'requests': 3, 'errors': 1, 'error_rate_pct': 33.3}), [])

    def test_error_rate_trips_above_threshold_with_real_volume(self):
        found = alerts._check_error_rate({'requests': 500, 'errors': 60, 'error_rate_pct': 12.0})
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0]['severity'], 'critical')

    def test_a_failing_check_reports_itself_instead_of_vanishing(self):
        def boom(_arg):
            raise RuntimeError('boom')
        boom.__name__ = '_check_latency'

        with patch.object(alerts, '_check_latency', boom):
            findings = alerts.evaluate()
        broken = [f for f in findings if f['key'] == 'check__check_latency']
        self.assertEqual(len(broken), 1)
        self.assertIn('RuntimeError: boom', broken[0]['detail'])


@override_settings(EMAIL_BACKEND=LOCMEM, C360_REPORT_RECIPIENTS=['ops@example.com'])
class DigestTests(TestCase):
    def setUp(self):
        mail.outbox = []
        now = timezone.now()
        MetricMinute.objects.create(minute=now - timedelta(minutes=5), instance='t:1',
                                    count=100, errors=2, client_errors=3, sum_ms=5000,
                                    p50_ms=40, p95_ms=120, p99_ms=300, max_ms=800)
        AuditEvent.objects.create(kind='api', method='GET', route='/api/boom/',
                                  status=500, duration_ms=900, username='rm.two')

    def test_daily_digest_sends_with_attachments(self):
        report = digest.send_daily()
        self.assertEqual(report['sent'], 1)
        message = mail.outbox[0]
        self.assertIn('Daily operations report', message.subject)
        names = [a[0] for a in message.attachments]
        self.assertTrue(any(n.endswith('.xlsx') for n in names))
        self.assertTrue(any(n.endswith('.csv') for n in names))
        # An HTML alternative AND a non-empty text body — a blank text part is what
        # gets a report filed as spam.
        self.assertTrue(message.body.strip())
        self.assertEqual(message.alternatives[0][1], 'text/html')

    def test_digest_reports_measured_numbers(self):
        report = digest.build(24 * 60, 'Daily operations report')
        self.assertEqual(report['summary']['requests'], 100)
        self.assertEqual(report['summary']['errors'], 2)
        self.assertIn('2.0%', report['html'])

    def test_digest_says_so_when_there_was_no_traffic(self):
        MetricMinute.objects.all().delete()
        report = digest.build(60, 'Daily operations report')
        self.assertIn('no traffic', report['subject'])
        self.assertIn('No requests were served', report['html'])


@override_settings(EMAIL_BACKEND=LOCMEM, C360_REPORT_RECIPIENTS=['ops@example.com'])
class ErrorLogMailTests(TestCase):
    def setUp(self):
        mail.outbox = []

    def test_silent_when_there_are_no_errors(self):
        """An hourly 'no errors' email is a mail rule waiting to happen."""
        self.assertIsNone(errorlog.build(60))
        result = errorlog.send(60)
        self.assertEqual(result['sent'], 0)
        self.assertEqual(len(mail.outbox), 0)

    def test_groups_repeats_by_endpoint(self):
        for _ in range(4):
            AuditEvent.objects.create(kind='api', method='GET', route='/api/boom/',
                                      status=500, username='rm.one')
        AuditEvent.objects.create(kind='api', method='GET', route='/api/nope/',
                                  status=404, username='rm.two')
        report = errorlog.build(60)
        self.assertIsNotNone(report)
        self.assertIn('4 server / 1 client', report['subject'])
        self.assertIn('/api/boom/', report['text'])
        self.assertEqual(errorlog.send(60)['sent'], 1)


class MailerGuardTests(TestCase):
    @override_settings(EMAIL_BACKEND=LOCMEM, C360_REPORT_RECIPIENTS=[])
    def test_no_recipients_is_a_logged_no_op_not_a_crash(self):
        from c360.reports import mailer
        self.assertEqual(mailer.send('s', '<p>h</p>', 't'), 0)

    @override_settings(EMAIL_BACKEND='c360.tests.test_reporting.ExplodingBackend',
                       C360_REPORT_RECIPIENTS=['ops@example.com'])
    def test_smtp_failure_never_propagates(self):
        """A cron'd monitoring mail that raises takes the next report down with it."""
        from c360.reports import mailer
        self.assertEqual(mailer.send('s', '<p>h</p>', 't'), 0)


class ExplodingBackend:
    """A mail backend that fails, to prove the mailer swallows and logs it."""

    def __init__(self, *args, **kwargs):
        raise RuntimeError('smtp is down')
